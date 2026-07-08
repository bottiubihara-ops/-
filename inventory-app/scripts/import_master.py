#!/usr/bin/env python3
"""実地棚卸表(.xlsm)から品目マスタを抽出し、棚卸しアプリ用ワークブック(inventory.xlsx)を生成する。

使い方:
    python3 import_master.py <実地棚卸表.xlsm> <出力先inventory.xlsx>

- 元ファイルはシートごとに「原料コード(A列) / 原料名(C列)」の印刷用フォーマットを想定
- ふりがな(rPh)は除去して品名を取り出す
- 出力ワークブックは3シート構成:
    品目マスタ  (テーブル名: MasterTable)
    棚卸し記録  (テーブル名: RecordTable)
    入力者      (テーブル名: MemberTable)
  → このファイルをそのままTeamsにアップロードすれば共有DBとして使える
"""
import sys
import re
import zipfile
import xml.etree.ElementTree as ET

M = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
R = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'

CODE_COL = 1   # A列: 原料コード
NAME_COL = 3   # C列: 原料名
NOTE_COL = 9   # I列: 型番（フィルターシートのみ）

# シートごとの単位（元帳票の記入単位に合わせる）
SHEET_UNITS = {
    '原料豆': 'kg',
    '副原料': 'kg',
    'フィルター': '箱/本',
}


def shared_strings(z: zipfile.ZipFile) -> list[str]:
    """sharedStrings.xml を読み、ふりがな(rPh)を除外した文字列リストを返す。"""
    try:
        root = ET.fromstring(z.read('xl/sharedStrings.xml'))
    except KeyError:
        return []
    out = []
    for si in root.findall(f'{M}si'):
        parts = []
        t = si.find(f'{M}t')
        if t is not None:
            parts.append(t.text or '')
        for r in si.findall(f'{M}r'):  # リッチテキストラン（rPhは含めない）
            rt = r.find(f'{M}t')
            if rt is not None:
                parts.append(rt.text or '')
        out.append(''.join(parts))
    return out


def col_index(ref: str) -> int:
    s = 0
    for ch in re.match(r'([A-Z]+)', ref).group(1):
        s = s * 26 + ord(ch) - 64
    return s


def sheet_cells(z: zipfile.ZipFile, path: str, sst: list[str]) -> dict:
    """{(row, col): 値} を返す。"""
    root = ET.fromstring(z.read(path))
    cells = {}
    for row in root.find(f'{M}sheetData').findall(f'{M}row'):
        rn = int(row.get('r'))
        for c in row.findall(f'{M}c'):
            t, v = c.get('t'), c.find(f'{M}v')
            if t == 's' and v is not None:
                val = sst[int(v.text)]
            elif t == 'inlineStr':
                val = ''.join(x.text or '' for x in c.iter(f'{M}t'))
            elif v is not None:
                val = v.text
            else:
                val = ''
            if val is not None and str(val).strip() != '':
                cells[(rn, col_index(c.get('r')))] = str(val)
    return cells


def normalize(text: str) -> str:
    """改行→スペース、連続空白を1つに。"""
    text = text.replace('\r', ' ').replace('\n', ' ')
    text = re.sub(r'[ 　]+', ' ', text)
    return text.strip()


def normalize_code(raw: str) -> str:
    """数値コードの小数点表記(300076000.0)を整数文字列へ。非数値はそのまま。"""
    raw = raw.strip()
    try:
        f = float(raw)
        if f == int(f):
            return str(int(f))
    except ValueError:
        pass
    return raw


def extract_items(z: zipfile.ZipFile) -> list[dict]:
    sst = shared_strings(z)
    wb = ET.fromstring(z.read('xl/workbook.xml'))
    rels = {rel.get('Id'): rel.get('Target')
            for rel in ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))}
    items = []
    for sheet in wb.findall(f'.//{M}sheet'):
        name = sheet.get('name')
        target = rels[sheet.get(f'{R}id')]
        path = 'xl/' + target if not target.startswith('xl/') else target
        cells = sheet_cells(z, path, sst)
        for (rn, cn), val in sorted(cells.items()):
            if cn != NAME_COL:
                continue
            item_name = normalize(val)
            if not item_name or '原料名' in item_name:  # ヘッダー行を除外
                continue
            if rn < 4:  # タイトル・凡例行を除外
                continue
            code = normalize_code(cells.get((rn, CODE_COL), ''))
            if '原料コード' in code:
                code = ''
            note = ''
            model_no = normalize(cells.get((rn, NOTE_COL), ''))
            if name == 'フィルター' and model_no:
                note = f'型番: {model_no}'
            if code == 'テスト':
                code, note = '', 'テスト品'
            items.append({
                'code': code,
                'name': item_name,
                'category': name,
                'unit': SHEET_UNITS.get(name, ''),
                'note': note,
            })
    return items


def build_workbook(items: list[dict], out_path: str) -> None:
    from openpyxl import Workbook
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    style = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)

    # --- 品目マスタ ---
    ws = wb.active
    ws.title = '品目マスタ'
    master_headers = ['品目ID', '原料コード', '品名', 'カテゴリ', '単位', '備考']
    ws.append(master_headers)
    for i, it in enumerate(items, start=1):
        ws.append([f'M{i:03d}', it['code'], it['name'], it['category'], it['unit'], it['note']])
    ref = f'A1:{get_column_letter(len(master_headers))}{len(items) + 1}'
    table = Table(displayName='MasterTable', ref=ref)
    table.tableStyleInfo = style
    ws.add_table(table)
    ws.column_dimensions['C'].width = 36
    for col in ('A', 'B', 'D', 'E', 'F'):
        ws.column_dimensions[col].width = 14

    # --- 棚卸し記録 ---
    ws = wb.create_sheet('棚卸し記録')
    record_headers = ['記録ID', '記録日時', '棚卸し月', '原料コード', '品名', 'カテゴリ',
                      '数量', '単位', '半端量', '保管場所', '期限日', '期限区分', '入力者', '備考']
    ws.append(record_headers)
    ref = f'A1:{get_column_letter(len(record_headers))}2'  # テーブルは最低1データ行が必要なため空行を含める
    table = Table(displayName='RecordTable', ref=ref)
    table.tableStyleInfo = style
    ws.add_table(table)
    for i in range(1, len(record_headers) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 16
    ws.column_dimensions['E'].width = 36

    # --- 入力者 ---
    ws = wb.create_sheet('入力者')
    ws.append(['名前'])
    for name in ('入力者A', '入力者B', '入力者C'):
        ws.append([name])
    table = Table(displayName='MemberTable', ref='A1:A4')
    table.tableStyleInfo = style
    ws.add_table(table)
    ws.column_dimensions['A'].width = 20

    wb.save(out_path)
    fix_table_rel_targets(out_path)


def fix_table_rel_targets(path: str) -> None:
    """openpyxlはテーブル参照を絶対パス(/xl/tables/…)で書くが、
    exceljs等は相対パス(../tables/…)を前提とするため書き換える(Excel互換の標準形)。"""
    import io
    buf = io.BytesIO()
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename.startswith('xl/worksheets/_rels/'):
                data = data.replace(b'Target="/xl/tables/', b'Target="../tables/')
            zout.writestr(info, data)
    with open(path, 'wb') as f:
        f.write(buf.getvalue())


def main() -> None:
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(1)
    src, dst = sys.argv[1], sys.argv[2]
    with zipfile.ZipFile(src) as z:
        items = extract_items(z)
    by_cat = {}
    for it in items:
        by_cat.setdefault(it['category'], []).append(it)
    for cat, lst in by_cat.items():
        with_code = sum(1 for x in lst if x['code'])
        print(f'{cat}: {len(lst)}品目（うち原料コードあり {with_code}件）')
    print(f'合計: {len(items)}品目')
    build_workbook(items, dst)
    print(f'出力: {dst}')


if __name__ == '__main__':
    main()
