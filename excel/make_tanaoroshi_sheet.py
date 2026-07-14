#!/usr/bin/env python3
"""Teams共同編集用の実地棚卸シートを生成する。

アプリを使わず、Teams上のExcelを直接編集して棚卸しする運用（案A）のためのテンプレート。
品目マスタ(inventory-app/data/inventory.xlsx)から生成する。

構成:
  サマリ    : 棚卸し月の設定・シート別の進捗カウンタ・使い方
  原料豆    : 38品目（原料コード順）
  副原料    : 87品目（原料コード順）
  フィルター: 14品目（数量は箱・本の2列）
  設定      : 入力者リスト（ドロップダウンの元）

特徴:
  - 期限日を入れると期限区分（期限切れ/来月期限/再来月期限）を数式で自動判定
    （基準はサマリ!B2の棚卸し月。月初入力のズレ問題に対応）
  - 入力者・チェック状態はドロップダウン
  - 未入力の数量セルは黄色、期限切れ行は赤系の条件付き書式
"""
import sys
from datetime import date

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

MASTER_PATH = sys.argv[1] if len(sys.argv) > 1 else "inventory-app/data/inventory.xlsx"
OUT_PATH = sys.argv[2] if len(sys.argv) > 2 else "excel/棚卸しシート_共同編集版.xlsx"

HEAD_FILL = PatternFill("solid", fgColor="DDE5F0")
TITLE_FONT = Font(bold=True, size=13)
NOTE_FONT = Font(size=9, color="777777")
HEAD_FONT = Font(bold=True, size=10)
THIN = Side(style="thin", color="C8CFD8")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
YELLOW = PatternFill("solid", fgColor="FFF3C4")
RED = PatternFill("solid", fgColor="FBE0E0")
ORANGE = PatternFill("solid", fgColor="FFE9D2")

DATA_START = 4  # 1:タイトル 2:注意書き 3:ヘッダー


def load_master():
    wb = load_workbook(MASTER_PATH)
    ws = wb["品目マスタ"]
    items = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[2]:
            continue
        items.append({
            "code": str(row[1]) if row[1] is not None else "",
            "name": str(row[2]),
            "cat": str(row[3] or ""),
            "unit": str(row[4] or "個"),
            "note": str(row[5]) if row[5] else "",
        })
    return items


def code_order(it):
    return (it["code"] == "", it["code"], it["name"])


def style_header(ws, row, cols):
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def add_member_dv(ws, col_letter, last_row):
    dv = DataValidation(type="list", formula1="=設定!$A$2:$A$11", allow_blank=True, showErrorMessage=False)
    ws.add_data_validation(dv)
    dv.add(f"{col_letter}{DATA_START}:{col_letter}{last_row}")


def add_check_dv(ws, col_letter, last_row):
    dv = DataValidation(type="list", formula1='"1回目完了,ダブルチェック完了"', allow_blank=True, showErrorMessage=False)
    ws.add_data_validation(dv)
    dv.add(f"{col_letter}{DATA_START}:{col_letter}{last_row}")


def build_ingredient_sheet(wb, name, items):
    """原料豆・副原料シート（12列・期限判定数式つき）"""
    ws = wb.create_sheet(name)
    last = DATA_START + len(items) - 1
    ws["A1"] = f"実地棚卸表（{name}）"
    ws["A1"].font = TITLE_FONT
    ws["D1"] = '="棚卸し月: "&TEXT(サマリ!$B$2,"yyyy年m月")'
    ws["D1"].font = Font(bold=True, color="2B5CB8")
    ws["A2"] = "①数量には期限切れ分を含めない ②期限日を入れると期限区分は自動判定（手動で上書き可） ③行の追加・削除・並び替えはしない"
    ws["A2"].font = NOTE_FONT
    headers = ["原料コード", "品名", "数量", "単位", "期限切れ量", "期限切れ日",
               "期限日", "期限区分(自動)", "期限量", "入力者", "チェック", "備考"]
    for i, h in enumerate(headers, start=1):
        ws.cell(row=3, column=i, value=h)
    style_header(ws, 3, len(headers))
    for r, it in enumerate(items, start=DATA_START):
        ws.cell(row=r, column=1, value=it["code"])
        ws.cell(row=r, column=2, value=it["name"])
        ws.cell(row=r, column=4, value="kg")
        # 期限区分: 棚卸し月(サマリ!B2)基準で自動判定
        ws.cell(row=r, column=8, value=(
            f'=IF($G{r}="","",IF(EDATE(サマリ!$B$2,1)>$G{r},"期限切れ",'
            f'IF(EDATE(サマリ!$B$2,2)>$G{r},"来月期限",'
            f'IF(EDATE(サマリ!$B$2,3)>$G{r},"再来月期限",""))))'
        ))
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = BORDER
            if c in (6, 7):
                cell.number_format = "yyyy-mm-dd"
    widths = [12, 34, 9, 6, 10, 12, 12, 12, 9, 10, 14, 14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "C4"
    add_member_dv(ws, "J", last)
    add_check_dv(ws, "K", last)
    rng = f"A{DATA_START}:L{last}"
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'$H{DATA_START}="期限切れ"'], fill=RED))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'$H{DATA_START}="来月期限"'], fill=ORANGE))
    ws.conditional_formatting.add(
        f"C{DATA_START}:C{last}",
        FormulaRule(formula=[f'ISBLANK($C{DATA_START})'], fill=YELLOW),
    )
    return last


def build_filter_sheet(wb, items):
    """フィルターシート（数量は箱・本の2列、期限列なし）"""
    ws = wb.create_sheet("フィルター")
    last = DATA_START + len(items) - 1
    ws["A1"] = "実地棚卸表（フィルター）"
    ws["A1"].font = TITLE_FONT
    ws["D1"] = '="棚卸し月: "&TEXT(サマリ!$B$2,"yyyy年m月")'
    ws["D1"].font = Font(bold=True, color="2B5CB8")
    ws["A2"] = "①箱・本それぞれの数を記入（無ければ0） ②行の追加・削除・並び替えはしない"
    ws["A2"].font = NOTE_FONT
    headers = ["原料コード", "品名", "型番", "数量(箱)", "数量(本)", "入力者", "チェック", "備考"]
    for i, h in enumerate(headers, start=1):
        ws.cell(row=3, column=i, value=h)
    style_header(ws, 3, len(headers))
    for r, it in enumerate(items, start=DATA_START):
        ws.cell(row=r, column=1, value=it["code"])
        ws.cell(row=r, column=2, value=it["name"])
        ws.cell(row=r, column=3, value=it["note"].replace("型番: ", ""))
        for c in range(1, len(headers) + 1):
            ws.cell(row=r, column=c).border = BORDER
    widths = [12, 30, 22, 9, 9, 10, 14, 14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "C4"
    add_member_dv(ws, "F", last)
    add_check_dv(ws, "G", last)
    ws.conditional_formatting.add(
        f"D{DATA_START}:E{last}",
        FormulaRule(formula=[f'AND(ISBLANK($D{DATA_START}),ISBLANK($E{DATA_START}))'], fill=YELLOW),
    )
    return last


def build_summary(wb, ranges):
    ws = wb["サマリ"]
    ws["A1"] = "棚卸しサマリ"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = "棚卸し月（月初の日付を入力）"
    ws["A2"].font = HEAD_FONT
    today = date.today()
    ws["B2"] = date(today.year, today.month, 1)
    ws["B2"].number_format = 'yyyy"年"m"月"'
    ws["B2"].font = Font(bold=True, size=12, color="2B5CB8")
    ws["C2"] = "← 例: 6月末の棚卸しなら 2026/6/1 と入力（期限判定の基準になります）"
    ws["C2"].font = NOTE_FONT

    headers = ["シート", "対象品目", "入力済み", "未入力", "1回目完了", "ダブル完了"]
    for i, h in enumerate(headers, start=1):
        ws.cell(row=4, column=i, value=h)
    style_header(ws, 4, len(headers))
    r = 5
    for name, last, qty_expr, check_col in ranges:
        n = last - DATA_START + 1
        ws.cell(row=r, column=1, value=name)
        ws.cell(row=r, column=2, value=n)
        ws.cell(row=r, column=3, value=qty_expr)
        ws.cell(row=r, column=4, value=f"=B{r}-C{r}")
        ws.cell(row=r, column=5, value=f'=COUNTIF({name}!${check_col}${DATA_START}:${check_col}${last},"1回目完了")')
        ws.cell(row=r, column=6, value=f'=COUNTIF({name}!${check_col}${DATA_START}:${check_col}${last},"ダブルチェック完了")')
        for c in range(1, 7):
            ws.cell(row=r, column=c).border = BORDER
        r += 1
    ws.cell(row=r, column=1, value="合計").font = HEAD_FONT
    for c in range(2, 7):
        col = get_column_letter(c)
        ws.cell(row=r, column=c, value=f"=SUM({col}5:{col}{r - 1})").font = HEAD_FONT
        ws.cell(row=r, column=c).border = BORDER
    ws.cell(row=r, column=1).border = BORDER

    tips = [
        "■ 使い方",
        "1. 各シートを開き、上から順に「数量」を記入する（在庫が無ければ 0）",
        "2. 期限切れ分は数量に含めず「期限切れ量」「期限切れ日」に記入",
        "3. 期限日を入れると「期限区分」が自動で入る（来月/再来月のときは「期限量」も記入）",
        "4. 入力したら「入力者」を選び、確認が済んだら「チェック」を選ぶ",
        "5. 未入力の数量セルは黄色で表示される。全員の入力状況はこのサマリで確認",
        "※ 行の追加・削除・並び替えはしないでください（数式とカウントが狂います）",
        "※ 複数人で同時に開いて編集してOK（自動保存・自動同期）",
    ]
    for i, t in enumerate(tips):
        cell = ws.cell(row=r + 2 + i, column=1, value=t)
        cell.font = HEAD_FONT if i == 0 else Font(size=10)
    for col, w in zip("ABCDEF", [14, 12, 10, 10, 12, 12]):
        ws.column_dimensions[col].width = w
    ws.column_dimensions["C"].width = 60 if False else 12


def build_settings(wb):
    ws = wb.create_sheet("設定")
    ws["A1"] = "入力者リスト（自由に書き換えOK・最大10人）"
    ws["A1"].font = HEAD_FONT
    for i, name in enumerate(["入力者A", "入力者B", "入力者C"], start=2):
        ws.cell(row=i, column=1, value=name)
    for r in range(2, 12):
        ws.cell(row=r, column=1).border = BORDER
    ws.column_dimensions["A"].width = 24


def main():
    items = load_master()
    beans = sorted([i for i in items if i["cat"] == "原料豆"], key=code_order)
    subs = sorted([i for i in items if i["cat"] == "副原料"], key=code_order)
    filters = [i for i in items if i["cat"] == "フィルター"]

    wb = Workbook()
    wb.active.title = "サマリ"
    last_b = build_ingredient_sheet(wb, "原料豆", beans)
    last_s = build_ingredient_sheet(wb, "副原料", subs)
    last_f = build_filter_sheet(wb, filters)
    build_settings(wb)
    build_summary(wb, [
        ("原料豆", last_b, f"=COUNT(原料豆!$C${DATA_START}:$C${last_b})", "K"),
        ("副原料", last_s, f"=COUNT(副原料!$C${DATA_START}:$C${last_s})", "K"),
        ("フィルター", last_f,
         # 箱・本のどちらかが数値なら入力済み（両方入力の重複分はCOUNTIFSで差し引く）
         f"=COUNT(フィルター!$D${DATA_START}:$D${last_f})+COUNT(フィルター!$E${DATA_START}:$E${last_f})"
         f'-COUNTIFS(フィルター!$D${DATA_START}:$D${last_f},">=0",フィルター!$E${DATA_START}:$E${last_f},">=0")',
         "G"),
    ])
    wb.save(OUT_PATH)
    print(f"生成: {OUT_PATH}（原料豆{len(beans)} / 副原料{len(subs)} / フィルター{len(filters)}品目）")


if __name__ == "__main__":
    main()
